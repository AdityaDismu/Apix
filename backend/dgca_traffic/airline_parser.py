
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from .config import AIRLINE_METRICS_JSON


AIRLINES = {
    "AIX CONNECT",
    "AIR INDIA",
    "AIR INDIA EXPRESS",
    "AKASA AIR",
    "ALLIANCE AIR",
    "BLUEDART",
    "FLY91",
    "FLYBIG",
    "INDIA ONE AIR",
    "INDIGO",
    "QUIKJET CARGO",
    "SPICEJET",
    "STAR AIR",
    "VISTARA",
}


def number(value: object) -> float | None:
    """
    Convert a DGCA Excel value into a number.

    '-' and blank cells represent unavailable/not-reported
    values and are therefore returned as None.
    """

    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    value = str(value).strip()

    if value in {"-", ""}:
        return None

    value = value.replace(",", "")

    try:
        return float(value)
    except ValueError:
        return None


def normalize_airline_name(
    value: object,
) -> str:

    if value is None:
        return ""

    return str(value).strip().upper()


def find_airline_column(
    worksheet,
) -> int:
    """
    Locate the column containing airline/operator names.

    The Excel export may have title rows above the table,
    so we search the worksheet rather than assuming a
    fixed row or column.
    """

    for row_number in range(
        1,
        min(worksheet.max_row, 30) + 1,
    ):
        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if value is None:
                continue

            text = normalize_airline_name(
                value
            )

            if text in AIRLINES:
                return column_number

    raise ValueError(
        "Could not locate the DGCA airline column "
        "in the Excel workbook."
    )


def parse_airline_excel(
    excel_path: Path,
) -> dict:

    if not excel_path.exists():
        raise FileNotFoundError(
            f"DGCA airline Excel file not found: "
            f"{excel_path}"
        )

    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    airline_column = find_airline_column(
        worksheet
    )

    metrics = []

    for row_number in range(
        1,
        worksheet.max_row + 1,
    ):

        airline_value = worksheet.cell(
            row=row_number,
            column=airline_column,
        ).value

        airline = normalize_airline_name(
            airline_value
        )

        if airline not in AIRLINES:
            continue

        numeric_values = []

        # Preserve the numerical values appearing
        # in the same Excel row.
        #
        # We deliberately do not assign these values
        # to specific operational fields because that
        # mapping has not been independently verified.

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):

            if column_number == airline_column:
                continue

            value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            parsed = number(value)

            if parsed is not None:
                numeric_values.append(
                    parsed
                )

        metrics.append(
            {
                "airline": airline,
                "source_row": row_number,
                "raw_numeric_values":
                    numeric_values,
            }
        )

    workbook.close()

    result = {
        "reference_period": "2024-25",
        "source":
            "DGCA Table 4.01",
        "airlines_found":
            len(metrics),
        "airlines": metrics,
    }

    AIRLINE_METRICS_JSON.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    AIRLINE_METRICS_JSON.write_text(
        json.dumps(
            result,
            indent=2,
        ),
        encoding="utf-8",
    )

    return result


# Backward-compatible function name.
#
# build_dataset.py may still call parse_airline_pdf().
# Keeping this wrapper means we don't need to change
# the builder immediately.

def parse_airline_pdf(
    pdf_path: Path,
) -> dict:

    return parse_airline_excel(
        pdf_path
    )

